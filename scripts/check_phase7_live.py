"""Real-provider acceptance on the bundled FAKE student sample only.

Starts its own backend8028/frontend5528, closes only its own processes. Requires
backend/.env and network access to Geoapify, Leaflet CDN and imagery. Local-only
artifacts contain the fake sample; no API credentials/request URLs are logged.
"""
from datetime import datetime, timezone
import json
import os
import subprocess
import sys
import time

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright, expect
from check_phase2_browser import ROOT, wait_ready

OUT = ROOT / ".local-tests/phase-7"


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    processes = []
    summary = {"checked_at_utc": datetime.now(timezone.utc).isoformat(), "source": "live Geoapify via backend; bundled fake30 sample", "passed": False}
    with (OUT / "live-servers.log").open("w", encoding="utf8") as log:
        try:
            env = {**os.environ, "CORS_ORIGINS": "http://127.0.0.1:5528"}
            processes.append(subprocess.Popen([sys.executable, "-B", "-m", "uvicorn", "app.main:app", "--app-dir", "backend", "--host", "127.0.0.1", "--port", "8028"], cwd=ROOT, env=env, stdout=log, stderr=log))
            processes.append(subprocess.Popen([sys.executable, "-B", "-m", "http.server", "5528", "--bind", "127.0.0.1", "--directory", "frontend"], cwd=ROOT, stdout=log, stderr=log))
            wait_ready("http://127.0.0.1:8028/api/health", processes[0])
            wait_ready("http://127.0.0.1:5528", processes[1])
            with sync_playwright() as p:
                browser = p.chromium.launch(channel="msedge", headless=True)
                context = browser.new_context(accept_downloads=True, viewport={"width": 1440, "height": 1000})
                context.route("**/js/api.js", lambda route: route.fulfill(content_type="text/javascript", body=(ROOT / "frontend/js/api.js").read_text(encoding="utf8").replace(":8001", ":8028")))
                page = context.new_page()
                page.set_default_timeout(15000)
                lookups, errors = [], []
                def observe(request):
                    assert "apikey=" not in request.url.lower()
                    if request.url.endswith("/geocode/retry"):
                        lookups.append(request.post_data_json)
                page.on("request", observe)
                page.on("pageerror", lambda error: errors.append(str(error)))
                page.goto("http://127.0.0.1:5528", wait_until="domcontentloaded")
                expect(page.locator("#connection-status")).to_have_attribute("data-state", "online")
                started = time.monotonic()
                page.locator("#student-file").set_input_files(str(ROOT / "data/samples/danh_sach_sinh_vien_bao_cao.xlsx"))
                page.locator("#import-button").click()
                expect(page.locator("#student-rows tr")).to_have_count(30)
                expect(page.locator("#geocode-all")).to_be_enabled(timeout=150000)
                rows = page.evaluate("async()=>(await import('/js/state.js')).state.students")
                summary.update(students=len(rows), unique_addresses=len({s.get("dia_chi_da_sua") or s["dia_chi_goc"] for s in rows}),
                    automatic=sum(s.get("location_method") == "automatic" for s in rows),
                    review=sum(not s["location_confirmed"] and s["geocode_status"] != "error" for s in rows),
                    errors=sum(s["geocode_status"] == "error" for s in rows), manual_confirmations=sum(s.get("location_method") == "manual" for s in rows),
                    initial_lookup_calls=len(lookups), import_to_geocode_seconds=round(time.monotonic() - started, 3))
                # Persist diagnostics before any assertion, distinguishing outages from review.
                places = {}
                for row in rows:
                    address = row.get("dia_chi_da_sua") or row["dia_chi_goc"]
                    places.setdefault(address, {"status": row["geocode_status"], "lat": row["lat"], "lon": row["lon"], "label": row["vi_tri_dai_dien"],
                        "candidates": row.get("candidates", [])})
                (OUT / "live-places.json").write_text(json.dumps(places, ensure_ascii=False, indent=2), encoding="utf8")
                assert summary["automatic"] == 30 and summary["initial_lookup_calls"] == 9, "Live sample incomplete; see live-summary.json / live-places.json"
                snapshot = page.evaluate("async()=>(await import('/js/state.js')).groupingPayload()")
                (OUT / "confirmed-snapshot.json").write_text(json.dumps(snapshot, ensure_ascii=False, indent=2) + "\n", encoding="utf8")
                page.locator("#generate-groups").click()
                expect(page.locator("#grouping-status")).to_have_attribute("data-state", "success", timeout=30000)
                result = page.evaluate("async()=>(await import('/js/state.js')).state.groupingResult")
                assert result["k"] == 6
                summary["group_metrics"] = result["overall_metrics"]
                page.locator("#location-filter").select_option("all")
                card = page.locator(".location-card").first
                original_address = card.locator(".location-address").input_value()
                card.locator(".location-address").fill("Địa chỉ kiểm thử không xác định PHASE7ZZZZ")
                expect(page.locator("#export-groups")).to_be_disabled()
                assert page.evaluate("async()=>{let s=(await import('/js/state.js')).state;return s.students[0].lat===null&&s.students[0].lon===null&&s.groupingResult===null}")
                card.locator(".location-retry").click()
                expect(page.locator("#geocode-all")).to_be_enabled(timeout=45000)
                expect(page.locator("#generate-groups")).to_be_disabled()
                card.locator(".location-address").fill(original_address)
                card.locator(".location-retry").click()
                expect(page.locator("#location-progress")).to_contain_text("30/30", timeout=45000)
                page.locator("#generate-groups").click()
                expect(page.locator("#grouping-status")).to_have_attribute("data-state", "success", timeout=30000)
                with page.expect_download() as download:
                    page.locator("#export-groups").click()
                download.value.save_as(OUT / "live-corrected.xlsx")
                book = load_workbook(OUT / "live-corrected.xlsx")
                assert book.sheetnames == ["Danh_sach_nhom", "Thong_ke_nhom"] and [s.max_row for s in book] == [31, 7]
                assert sorted(r[2] for r in list(book.worksheets[0].values)[1:]) == sorted(s["ho_ten"] for s in rows)
                book.close()
                summary.update(correction_passed=True, total_lookup_calls=len(lookups))
                try:
                    expect(page.locator("#map-status")).to_have_attribute("data-state", "ready", timeout=45000)
                except AssertionError:
                    if page.locator("#retry-map").is_visible():
                        page.locator("#retry-map").click()
                    expect(page.locator("#map-status")).to_have_attribute("data-state", "ready", timeout=45000)
                assert page.locator(".leaflet-tile").evaluate_all("images=>images.length>0&&images.every(i=>i.complete&&i.naturalWidth>0)")
                summary["map_status"] = page.locator("#map-status").inner_text()
                summary["attribution"] = page.locator(".leaflet-control-attribution").inner_text()
                page.locator("#map").scroll_into_view_if_needed()
                page.screenshot(path=str(OUT / "live-desktop.png"))
                page.set_viewport_size({"width": 390, "height": 844})
                assert page.evaluate("document.documentElement.scrollWidth<=innerWidth")
                page.locator("#map").scroll_into_view_if_needed()
                page.screenshot(path=str(OUT / "live-mobile.png"))
                assert not errors, errors
                summary.update(passed=True, page_errors=errors)
                browser.close()
                print("PASS: live auto30/30, 9 initial lookups, address correction, 6groups, workbook30, real imagery, desktop/mobile", flush=True)
        finally:
            (OUT / "live-summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf8")
            for process in reversed(processes):
                if process.poll() is None:
                    process.terminate()
                    process.wait(timeout=10)


if __name__ == "__main__":
    main()

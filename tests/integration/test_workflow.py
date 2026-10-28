"""Exercise actual import/geocoding/grouping/export; only external HTTP is mocked."""
from io import BytesIO
from pathlib import Path

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.schemas.group import VerifiedStudent
from phase7_provider import make_app, provider

SAMPLE = Path(__file__).resolve().parents[2] / "data/samples/danh_sach_sinh_vien_bao_cao.xlsx"


def test_sample_preserves_identity_from_excel_to_export():
    with TestClient(make_app()) as client:
        imported = client.post("/api/students/import", files={"file": (SAMPLE.name, SAMPLE.read_bytes())})
        assert imported.status_code == 200
        rows = imported.json()["students"]
        book = load_workbook(SAMPLE)
        original = list(book.worksheets[0].values)[1:]
        book.close()
        assert len(rows) == len(original) == 30
        assert len({row["student_id"] for row in rows}) == 30
        confirmed = []
        for row, source in zip(rows, original):
            assert (row["ho_ten"], row["gioi_tinh"], row["dia_chi_goc"]) == source[:3]
            lookup = client.post("/api/geocode/retry", json={"student_id": row["student_id"], "address": row["dia_chi_da_sua"] or row["dia_chi_goc"]})
            assert lookup.status_code == 200
            reply = lookup.json()
            assert reply["student_id"] == row["student_id"]
            assert reply["status"] == "needs_review"
            assert len(reply["candidates"]) == 1
            candidate = reply["candidates"][0]
            assert candidate["eligible"] and candidate["auto_selectable"]
            assert not candidate["allowed_genders"] or row["gioi_tinh"] in candidate["allowed_genders"]
            row.update(lat=candidate["lat"], lon=candidate["lon"], vi_tri_dai_dien=candidate["formatted"],
                       geocode_status="verified", location_confirmed=True)
            confirmed.append({key: value for key, value in row.items() if key in VerifiedStudent.model_fields})
        request = {"students": confirmed, "target_male": 3, "target_female": 2}
        grouped = client.post("/api/groups/generate", json=request)
        assert grouped.status_code == 200
        result = grouped.json()
        assert result["k"] == 6
        members = [member for group in result["groups"] for member in group["members"]]
        assert sorted(members, key=lambda m: m["student_id"]) == sorted(confirmed, key=lambda m: m["student_id"])
        assert all((g["size"], g["male"], g["female"]) == (5, 3, 2) for g in result["groups"])
        exported = client.post("/api/groups/export", json=result)
        assert exported.status_code == 200
        book = load_workbook(BytesIO(exported.content))
        assert book.sheetnames == ["Danh_sach_nhom", "Thong_ke_nhom"]
        assert [s.max_row for s in book] == [31, 7]
        assert sorted(row[2] for row in list(book.worksheets[0].values)[1:]) == sorted(row["ho_ten"] for row in rows)
        book.close()


@pytest.mark.parametrize("status,code", [(429, "rate_limit"), (403, "credentials"), (503, "provider_error")])
def test_external_failure_is_not_reported_as_wrong_address_and_retry_recovers(status, code):
    failed = True
    def handler(request):
        return httpx.Response(status, text="phase7-fixture-secret") if failed else provider(request)
    application = make_app(handler)
    with TestClient(application) as client:
        data = {"student_id": "S1", "address": "Cư xá A, KTX Nông Lâm, phường Linh Xuân, TP.HCM"}
        response = client.post("/api/geocode/retry", json=data)
        assert response.json()["error_code"] == code
        assert response.json()["status"] == "error"
        assert response.json()["candidates"] == []
        assert "phase7-fixture-secret" not in response.text
        failed = False
        # Simulate cooldown expiry, no sleep or retry burst.
        application.state.geocoding.cooldown_until = 0
        response = client.post("/api/geocode/retry", json=data)
        assert response.json()["candidates"][0]["auto_selectable"]

import copy
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.schemas.export import ExportGroupsRequest
from app.schemas.group import GenerateGroupsRequest
from app.services.export_service import export_workbook
from app.services.grouping_service import generate_groups
from grouping_fixtures import payload


@pytest.fixture
def result():
    return generate_groups(GenerateGroupsRequest.model_validate(payload())).model_dump()


def test_workbook_has_exact_sheets_all_members_and_matching_statistics(result):
    book = load_workbook(BytesIO(export_workbook(ExportGroupsRequest.model_validate(result))))
    assert book.sheetnames == ["Danh_sach_nhom", "Thong_ke_nhom"]
    members = list(book.worksheets[0].values)
    stats = list(book.worksheets[1].values)
    assert members[0] == ("STT", "Nhóm", "Họ và tên", "Giới tính", "Nơi ở để học tập")
    assert len(members) == 31 and len(stats) == 7
    assert [row[0] for row in members[1:]] == list(range(1, 31))
    assert all(row[4] == "Điểm tổng hợp" for row in members[1:])
    assert len(members[0]) == 5  # no raw lat/lon columns
    for actual, group in zip(stats[1:], result["groups"]):
        assert actual[:4] == (group["group_id"], 5, 3, 2)
        assert actual[4] == pytest.approx(group["metrics"]["mean_distance_m"])
        assert actual[5] == pytest.approx(group["metrics"]["max_distance_m"])
    book.close()


def test_sorting_is_stable_and_input_is_not_mutated(result):
    shuffled = copy.deepcopy(result)
    shuffled["groups"].reverse()
    for group in shuffled["groups"]:
        group["members"].reverse()
    before = copy.deepcopy(shuffled)
    books = [load_workbook(BytesIO(export_workbook(ExportGroupsRequest.model_validate(data)))) for data in (result, shuffled)]
    assert list(books[0].worksheets[0].values) == list(books[1].worksheets[0].values)
    assert list(books[0].worksheets[1].values) == list(books[1].worksheets[1].values)
    assert shuffled == before
    for book in books:
        book.close()


@pytest.mark.parametrize("text", ['=HYPERLINK("https://example.invalid","click")', "+cmd", "-1+2", "@SUM(1,2)"])
def test_user_strings_are_literal_text_not_formulas_or_links(result, text):
    result["groups"][0]["members"][0].update(ho_ten=text, vi_tri_dai_dien=text)
    book = load_workbook(BytesIO(export_workbook(ExportGroupsRequest.model_validate(result))), data_only=False)
    row = list(book.worksheets[0].rows)[1]
    assert row[2].value == text and row[4].value == text
    assert row[2].data_type == row[4].data_type == "s"
    assert row[2].hyperlink is None
    book.close()


@pytest.mark.parametrize("mutation", [
    "duplicate_member", "missing_assignment", "wrong_assignment", "duplicate_group", "size", "gender",
    "delta", "k", "target", "group_metric", "overall_metric", "pair_count", "unbalanced", "xml_text",
])
def test_inconsistent_export_is_rejected(result, mutation):
    first, second = result["groups"][:2]
    if mutation == "duplicate_member":
        first["members"][0] = copy.deepcopy(second["members"][0])
    elif mutation == "missing_assignment":
        result["assignments"].pop()
    elif mutation == "wrong_assignment":
        result["assignments"][0]["group_id"] = 999
    elif mutation == "duplicate_group":
        second["group_id"] = first["group_id"]
    elif mutation == "size":
        first["size"] = 6
    elif mutation == "gender":
        first["male"] = 4
    elif mutation == "delta":
        first["female_delta"] = 3
    elif mutation == "k":
        result["k"] = 7
    elif mutation == "target":
        result["target_male"] = -1
    elif mutation == "group_metric":
        first["metrics"]["mean_distance_m"] += 100
    elif mutation == "overall_metric":
        result["overall_metrics"]["max_distance_m"] += 100
    elif mutation == "pair_count":
        result["overall_metrics"]["pair_count"] += 1
    elif mutation == "unbalanced":
        first["members"].append(second["members"].pop())
    else:
        first["members"][0]["ho_ten"] = "Tên\x00không hợp lệ"
    with pytest.raises(ValueError):
        export_workbook(ExportGroupsRequest.model_validate(result))


def test_warning_notes_and_singleton_zero_distances():
    for data in (payload(12, 6), payload(1, 1, 1, 0)):
        result = generate_groups(GenerateGroupsRequest.model_validate(data)).model_dump()
        book = load_workbook(BytesIO(export_workbook(ExportGroupsRequest.model_validate(result))))
        row = list(book.worksheets[1].values)[1]
        if len(data["students"]) == 12:
            assert "3 Nam" in row[6] and "2 Nữ" in row[6] and "+1 Nữ" in row[6]
        else:
            assert row[4:6] == (0, 0)
        book.close()

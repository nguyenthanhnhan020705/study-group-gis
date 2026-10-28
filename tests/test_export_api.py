import copy
import json
from dataclasses import replace
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.core.config import Settings
from app.main import create_app
from grouping_fixtures import payload


@pytest.fixture
def api_result():
    app = create_app(Settings(cors_origins=("http://127.0.0.1:5500",)))
    with TestClient(app, raise_server_exceptions=False) as client:
        result = client.post("/api/groups/generate", json=payload()).json()
        yield client, result, app


def test_export_endpoint_returns_downloadable_workbook_with_cors(api_result):
    client, result, _ = api_result
    response = client.post("/api/groups/export", json=result, headers={"Origin": "http://127.0.0.1:5500"})
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert 'filename="danh_sach_nhom.xlsx"' in response.headers["content-disposition"]
    assert response.headers["access-control-allow-origin"] == "http://127.0.0.1:5500"
    assert response.headers["cache-control"] == "no-store"
    book = load_workbook(BytesIO(response.content))
    assert book.worksheets[0].max_row == 31 and book.worksheets[1].max_row == 7
    book.close()


def test_export_errors_do_not_echo_private_payload_or_nonfinite_values(api_result):
    client, result, _ = api_result
    for change in ({"lat": float("nan")}, {"ho_ten": ""}, {"location_confirmed": False}):
        data = copy.deepcopy(result)
        data["groups"][0]["members"][0].update(change)
        response = client.post("/api/groups/export", content=json.dumps(data), headers={"Content-Type": "application/json"})
        assert response.status_code == 422
        assert "Sinh viên tổng hợp" not in response.text and "NaN" not in response.text
    data = copy.deepcopy(result)
    data["groups"][0]["size"] += 1
    response = client.post("/api/groups/export", json=data)
    assert response.status_code == 422 and "Sinh viên tổng hợp" not in response.text


def test_export_honors_row_and_body_limits(api_result):
    client, result, app = api_result
    app.state.settings = replace(app.state.settings, max_student_rows=29)
    assert client.post("/api/groups/export", json=result).status_code == 413
    response = client.post("/api/groups/export", content=b" " * (8 * 1024 * 1024 + 1), headers={"Content-Type": "application/json"})
    assert response.status_code == 413


def test_export_has_no_server_history_or_file_and_does_not_regenerate_groups(api_result, monkeypatch):
    client, result, _ = api_result
    def forbidden(*args, **kwargs):
        raise AssertionError("Export must not regroup the user's current result")
    monkeypatch.setattr("app.api.groups.generate_groups", forbidden)
    assert client.post("/api/groups/export", json=result).status_code == 200
    assert client.get("/api/groups/export").status_code == 405

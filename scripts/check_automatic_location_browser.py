"""100-row actual import + synthetic provider responses, real Leaflet and grouping/export."""
import json
import os
import subprocess
import sys
from io import BytesIO

from openpyxl import Workbook, load_workbook
from playwright.sync_api import expect, sync_playwright
from check_phase2_browser import ROOT, wait_ready

OUT = ROOT / '.local-tests/automatic-location'
NLU = 'KTX Nông Lâm, phường Linh Xuân, TP.HCM'


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    book = Workbook(); book.active.append(['Họ và tên', 'Giới tính', 'Nơi ở để học tập', 'Cư xá/Khu vực'])
    for i in range(100):
        book.active.append([f'Sinh viên thử {i+1}', 'Nữ' if i == 99 else 'Nam', NLU, '' if i == 98 else 'Cư xá A'])
    path = OUT / '100-students.xlsx'; book.save(path); book.close()
    processes = []
    with (OUT / 'servers.log').open('w') as log:
        try:
            env = {**os.environ, 'CORS_ORIGINS':'http://127.0.0.1:5519'}
            processes.append(subprocess.Popen([sys.executable,'-m','uvicorn','app.main:app','--app-dir','backend','--port','8017'],cwd=ROOT,env=env,stdout=log,stderr=log))
            processes.append(subprocess.Popen([sys.executable,'-m','http.server','5519','--bind','127.0.0.1','--directory','frontend'],cwd=ROOT,stdout=log,stderr=log))
            wait_ready('http://127.0.0.1:8017/api/health',processes[0]); wait_ready('http://127.0.0.1:5519',processes[1])
            with sync_playwright() as p:
                browser=p.chromium.launch(channel='msedge',headless=True)
                page=browser.new_page(accept_downloads=True,viewport={'width':1440,'height':1000}); page.set_default_timeout(15000)
                errors=[]; calls=[]
                page.on('pageerror',lambda error:errors.append(str(error)))
                page.route('**/js/api.js',lambda route:route.fulfill(content_type='text/javascript',body=(ROOT/'frontend/js/api.js').read_text(encoding='utf8').replace(':8001',':8017')))
                def leaflet(route):
                    suffix='.css' if route.request.url.endswith('.css') else '.js'
                    route.fulfill(path=str(ROOT/f'.local-tests/phase-5/leaflet-1.9.4{suffix}'),content_type='text/css' if suffix=='.css' else 'text/javascript')
                page.route('https://unpkg.com/**',leaflet)
                page.route('https://tile.openstreetmap.org/**',lambda route:route.abort())
                page.route('**/api/map/tiles/**',lambda route:route.abort())
                def result(data):
                    explicit='Cư xá A' in data['address']
                    return {**data,'status':'needs_review','message':'Synthetic provider boundary',
                            'candidates':[{'lat':10.871312,'lon':106.7911041,'formatted':f'{NLU} — Cư xá A',
                                           'result_type':'amenity','confidence':None,'eligible':True,'source':'geoapify_places',
                                           'auto_selectable':explicit,'allowed_genders':['Nam']}]}
                def respond(route):
                    data=route.request.post_data_json;calls.append(data)
                    route.fulfill(json=result(data))
                page.route('**/api/geocode/retry',respond)
                page.goto('http://127.0.0.1:5519',wait_until='domcontentloaded')
                page.locator('#student-file').set_input_files(str(path));page.locator('#import-button').click()
                expect(page.locator('#location-progress')).to_contain_text('98/100',timeout=30000)
                expect(page.locator('#geocode-all')).to_be_enabled()
                assert len(calls)==2, len(calls)
                assert page.locator('.location-card:visible').count()==2
                assert page.evaluate("async()=>{const s=(await import('/js/state.js')).state;return s.students.filter(x=>x.location_method==='automatic').length===98&&!s.students[98].location_confirmed&&!s.students[99].location_confirmed}")
                assert page.locator('.leaflet-interactive').count()==1
                page.locator('#location-filter').select_option('all')
                assert page.locator('.location-card:visible').count()==100
                page.locator('#location-filter').select_option('review')
                page.screenshot(path=str(OUT/'exceptions.png'))
                print('PASS: actual100-row import auto-starts;98 verified,2 exceptions;2 lookup calls;real Leaflet marker',flush=True)

                # Delayed automatic match cannot confirm an edited row.
                pending=[];page.unroute('**/api/geocode/retry');page.route('**/api/geocode/retry',lambda route:pending.append(route))
                card=page.locator('.location-card').nth(98)
                card.locator('.location-retry').click()
                page.wait_for_timeout(100);assert len(pending)==1
                data=pending[0].request.post_data_json
                card.locator('.location-address').fill('Địa chỉ khác chưa đủ thông tin')
                pending[0].fulfill(json=result({**data,'address':'Cư xá A, '+NLU}))
                expect(card.locator('.location-confirm')).to_be_disabled()
                assert page.evaluate("async()=>!(await import('/js/state.js')).state.students[98].location_confirmed")
                print('PASS: stale result cannot auto-confirm edited address',flush=True)

                # New import cancels old work and does not accept late results.
                card.locator('.location-retry').click();page.wait_for_timeout(100)
                page.locator('#student-file').set_input_files({'name':'new-list.xlsx','mimeType':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','buffer':path.read_bytes()})
                expect(page.locator('.location-card')).to_have_count(0)
                try: pending[-1].fulfill(json=result(data))
                except Exception: pass
                assert page.evaluate("async()=>(await import('/js/state.js')).state.students.length===0")
                print('PASS: new file clears automatic queue and ignores old results',flush=True)
                page.unroute('**/api/geocode/retry')
                def complete(route):
                    data=route.request.post_data_json; reply=result(data)
                    female='Cư xá B' in data['address']
                    reply['candidates'][0].update(auto_selectable=True,allowed_genders=['Nữ' if female else 'Nam'],
                                                   formatted=f'{NLU} — Cư xá '+('B' if female else 'A'))
                    route.fulfill(json=reply)
                page.route('**/api/geocode/retry',complete)
                book=Workbook();book.active.append(['Họ và tên','Giới tính','Nơi ở để học tập','Cư xá/Khu vực'])
                for i in range(5):book.active.append([f'Sinh viên tích hợp{i}', 'Nam' if i<3 else 'Nữ',NLU,'Cư xá A' if i<3 else 'Cư xá B'])
                buffer=BytesIO();book.save(buffer);book.close()
                page.locator('#student-file').set_input_files({'name':'five.xlsx','mimeType':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','buffer':buffer.getvalue()})
                page.locator('#import-button').click()
                expect(page.locator('#location-progress')).to_contain_text('5/5')
                expect(page.locator('#generate-groups')).to_be_enabled()
                page.locator('#generate-groups').click()
                expect(page.locator('#grouping-status')).to_have_attribute('data-state','success')
                with page.expect_download() as pending_download:page.locator('#export-groups').click()
                pending_download.value.save_as(OUT/'automatic-groups.xlsx')
                book=load_workbook(OUT/'automatic-groups.xlsx')
                assert [sheet.max_row for sheet in book]==[6,2];book.close()
                print('PASS: automatic matches flow through actual grouping API and downloadable two-sheet Excel',flush=True)
                page.set_viewport_size({'width':390,'height':844})
                assert page.evaluate('document.documentElement.scrollWidth<=innerWidth')
                page.locator('#location-title').scroll_into_view_if_needed();page.screenshot(path=str(OUT/'mobile.png'))
                page.unroute('**/api/geocode/retry');failures=[]
                def unavailable(route):
                    failures.append(route.request.post_data_json)
                    route.fulfill(json={**failures[-1],'status':'error','error_code':'provider_error','message':'Provider unavailable','candidates':[]})
                page.route('**/api/geocode/retry',unavailable)
                page.locator('#student-file').set_input_files(str(path));page.locator('#import-button').click()
                expect(page.locator('#location-progress')).to_contain_text('Đã tạm dừng')
                expect(page.locator('#geocode-all')).to_be_enabled()
                assert len(failures)==1
                print('PASS: backend provider outage pauses100-row queue after one failure with retry available',flush=True)
                assert not errors,errors
                browser.close()
        finally:
            for process in reversed(processes):
                if process.poll() is None:process.terminate();process.wait(timeout=10)


if __name__=='__main__':main()

"""Bounded Excel import, retaining row identity and all invalid records."""

from io import BytesIO
from zipfile import ZipFile

from openpyxl import load_workbook

from app.schemas.student import ImportErrorDetail, ImportResult, ImportSummary, Student
from app.services.known_places import normalize

REQUIRED_COLUMNS = ("Họ và tên", "Giới tính", "Nơi ở để học tập")


class ImportProblem(Exception):
    def __init__(self, message: str, status_code: int = 422, *, field: str = "file", excel_row: int | None = None):
        super().__init__(message)
        self.status_code = status_code
        self.errors = [ImportErrorDetail(field=field, excel_row=excel_row, message=message)]


def parse_students(content: bytes, *, max_rows: int = 1000, max_uncompressed_bytes: int = 20 * 1024 * 1024) -> ImportResult:
    workbook = None
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(entry.file_size for entry in entries) > max_uncompressed_bytes:
                raise ImportProblem("File vượt giới hạn nội dung giải nén cho phép.", 413)
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        if not workbook.worksheets:
            raise ImportProblem("File không có sheet dữ liệu.")
        sheet = workbook.worksheets[0]
        # Do not trust a producer's incorrect used-range metadata.
        sheet.reset_dimensions()
        rows = sheet.iter_rows()
        first_row = next(rows, ())
        if len(first_row) > 256:
            raise ImportProblem("Sheet vượt giới hạn 256 cột.")
        headers = [str(cell.value).strip() if cell.value is not None else "" for cell in first_row]
        for column in REQUIRED_COLUMNS:
            if headers.count(column) != 1:
                message = f"Thiếu cột bắt buộc: {column}." if column not in headers else f"Cột bắt buộc bị trùng: {column}."
                raise ImportProblem(message, field=column, excel_row=1)
        indices = [headers.index(column) for column in REQUIRED_COLUMNS]
        hint_names = ("Cư xá/Khu vực", "Điểm cần xác nhận")
        hint_columns = [(name, headers.index(name)) for name in hint_names if name in headers]
        if any(headers.count(name) > 1 for name in hint_names):
            raise ImportProblem("Cột địa điểm bổ sung bị trùng.")
        data_rows = []
        for excel_row, row in enumerate(rows, start=2):
            if excel_row > max_rows + 1:
                raise ImportProblem(f"Sheet vượt giới hạn {max_rows} dòng dữ liệu (không tính tiêu đề).")
            if len(row) > 256:
                raise ImportProblem("Sheet vượt giới hạn 256 cột.", excel_row=excel_row)
            data_rows.append((excel_row, row))
        while data_rows and all(cell.value is None for cell in data_rows[-1][1]):
            data_rows.pop()
        if not data_rows:
            raise ImportProblem("File không có sinh viên để nhập.")

        students, errors = [], []
        for excel_row, row in data_rows:
            student_id = f"S{excel_row:06d}"
            values = []
            for column, index in zip(REQUIRED_COLUMNS, indices):
                cell = row[index] if index < len(row) else None
                raw = cell.value if cell else None
                value = str(raw).strip() if raw is not None else ""
                message = None
                if not value:
                    message = "Không được để trống."
                elif cell.data_type == "f":
                    message = "Phải nhập văn bản, không dùng công thức Excel."
                elif cell.data_type == "e":
                    message = "Ô đang chứa lỗi Excel. Hãy sửa thành văn bản hợp lệ."
                elif not isinstance(raw, str):
                    message = "Giá trị phải là chuỗi văn bản."
                elif len(value) > 2000:
                    message = "Giá trị vượt giới hạn 2.000 ký tự."
                elif column == "Giới tính" and value not in ("Nam", "Nữ"):
                    message = "Chỉ chấp nhận Nam hoặc Nữ (phân biệt chữ hoa và dấu)."
                if message:
                    errors.append(ImportErrorDetail(excel_row=excel_row, student_id=student_id, field=column, message=message))
                values.append(" ".join(value.split()) if column == "Họ và tên" else value)
            hints = []
            for name, index in hint_columns:
                cell = row[index] if index < len(row) else None
                raw = cell.value if cell else None
                if raw is None or raw == "":
                    continue
                if cell.data_type in {"f", "e"} or not isinstance(raw, str) or len(raw.strip()) > 2000:
                    errors.append(ImportErrorDetail(excel_row=excel_row, student_id=student_id, field=name,
                                                   message="Địa điểm bổ sung phải là văn bản tối đa 2.000 ký tự, không dùng công thức/lỗi Excel."))
                elif raw.strip() and normalize(raw).replace("ktx", "ky tuc xa") not in normalize(values[2]).replace("ktx", "ky tuc xa") and raw.strip() not in hints:
                    hints.append(raw.strip())
            lookup_address = ", ".join([*hints, values[2]]) if hints else None
            if lookup_address and len(lookup_address) > 2000:
                errors.append(ImportErrorDetail(excel_row=excel_row, student_id=student_id, field="Nơi ở để học tập",
                                               message="Địa chỉ kèm địa điểm bổ sung vượt 2.000 ký tự."))
            students.append(Student(student_id=student_id, excel_row=excel_row, ho_ten=values[0], gioi_tinh=values[1], dia_chi_goc=values[2], dia_chi_da_sua=lookup_address))
        summary = ImportSummary(
            total=len(students),
            male=sum(student.gioi_tinh == "Nam" for student in students),
            female=sum(student.gioi_tinh == "Nữ" for student in students),
            invalid_rows=len({error.excel_row for error in errors}),
        )
        return ImportResult(students=students, errors=errors, summary=summary, can_proceed=not errors, worksheet=sheet.title)
    except ImportProblem:
        raise
    except Exception as exc:
        # Parser exception details may contain workbook contents; never expose them.
        raise ImportProblem("Không đọc được file Excel. Hãy kiểm tra file .xlsx và tải lại.", 400) from exc
    finally:
        if workbook is not None:
            workbook.close()

"""Offline end-to-end: actual APIs/UI, synthetic external Geoapify transport.

Uses own ports 8027/5527. No provider credentials or network needed. Leaflet
is blocked here; actual imagery is checked separately by check_phase7_live.py.
"""
import json
import subprocess
import sys
from io import BytesIO

from openpyxl import Workbook, load_workbook
from playwright.sync_api import expect, sync_playwright
from check_phase2_browser import ROOT, wait_ready

OUT = ROOT / ".local-tests/phase-7"
NLU = "KTX Nông Lâm, phường Linh Xuân, TP.HCM"


def upload(page, rows):
    book = Workbook()
    book.active.append(["Họ và tên", "Giới tính", "Nơi ở để học tập", "Cư xá/Khu vực"])
    for row in rows:
        book.active.append(row)
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    page.locator("#student-file").set_input_files({"name": "synthetic.xlsx", "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "buffer": buffer.getvalue()})
    page.locator("#import-button").click()
    expect(page.locator("#student-rows tr")).to_have_count(len(rows))
    expect(page.locator("#geocode-all")).to_be_enabled(timeout=30000)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    processes = []
    with (OUT / "offline-servers.log").open("w", encoding="utf8") as log:
        try:
            processes.append(subprocess.Popen([sys.executable, "-B", "-m", "uvicorn", "phase7_provider:app", "--app-dir", "tests", "--host", "127.0.0.1", "--port", "8027"], cwd=ROOT, stdout=log, stderr=log))
            processes.append(subprocess.Popen([sys.executable, "-B", "-m", "http.server", "5527", "--bind", "127.0.0.1", "--directory", "frontend"], cwd=ROOT, stdout=log, stderr=log))
            wait_ready("http://127.0.0.1:8027/api/health", processes[0])
            wait_ready("http://127.0.0.1:5527", processes[1])
            with sync_playwright() as p:
                browser = p.chromium.launch(channel="msedge", headless=True)
                contexts = [browser.new_context(accept_downloads=True, viewport={"width": 1440, "height": 1000}) for _ in range(2)]
                pages = []
                errors = []
                for context in contexts:
                    def route_request(route):
                        url = route.request.url
                        if url.endswith("/js/api.js"):
                            route.fulfill(content_type="text/javascript", body=(ROOT / "frontend/js/api.js").read_text(encoding="utf8").replace(":8001", ":8027"))
                        elif url.startswith(("http://127.0.0.1:5527/", "http://127.0.0.1:8027/")):
                            route.continue_()
                        else:
                            route.abort()
                    context.route("**/*", route_request)
                    page = context.new_page()
                    page.on("pageerror", lambda error: errors.append(str(error)))
                    page.set_default_timeout(15000)
                    page.goto("http://127.0.0.1:5527", wait_until="domcontentloaded")
                    pages.append(page)
                first, second = pages
                calls = []
                first.on("request", lambda request: calls.append(request.post_data_json) if request.url.endswith("/geocode/retry") else None)
                first.locator("#student-file").set_input_files(str(ROOT / "data/samples/danh_sach_sinh_vien_bao_cao.xlsx"))
                first.locator("#import-button").click()
                expect(first.locator("#location-progress")).to_contain_text("30/30", timeout=30000)
                assert len(calls) == 9
                before = first.evaluate("async()=>(await import('/js/state.js')).state.students.map(s=>[s.student_id,s.ho_ten,s.gioi_tinh,s.dia_chi_goc])")
                first.locator("#generate-groups").click()
                expect(first.locator("#grouping-status")).to_have_attribute("data-state", "success")
                with first.expect_download() as download:
                    first.locator("#export-groups").click()
                download.value.save_as(OUT / "offline-groups.xlsx")
                book = load_workbook(OUT / "offline-groups.xlsx")
                assert book.sheetnames == ["Danh_sach_nhom", "Thong_ke_nhom"]
                assert [s.max_row for s in book] == [31, 7]
                book.close()
                print("PASS: Excel30 -> real geocoding boundary9 -> auto30 -> groups6 -> workbook30", flush=True)

                upload(second, [["Phiên riêng Nam", "Nam", NLU, "Cư xá A"], ["Phiên riêng Nữ", "Nữ", NLU, "Cư xá A"]])
                expect(second.locator("#location-progress")).to_contain_text("1/2")
                expect(second.locator("#generate-groups")).to_be_disabled()
                assert second.evaluate("async()=>{let s=(await import('/js/state.js')).state;return s.students[0].location_confirmed&&!s.students[1].location_confirmed}")
                assert first.evaluate("async()=>(await import('/js/state.js')).state.students.map(s=>[s.student_id,s.ho_ten,s.gioi_tinh,s.dia_chi_goc])") == before
                # Explicit retry bypasses cache and resolves the incompatible gender.
                card = second.locator(".location-card").nth(1)
                card.locator(".location-address").fill("Cư xá B, " + NLU)
                card.locator(".location-retry").click()
                expect(second.locator("#location-progress")).to_contain_text("2/2")
                second.locator("#generate-groups").click()
                expect(second.locator("#grouping-status")).to_have_attribute("data-state", "success")
                assert second.evaluate("async()=>(await import('/js/state.js')).state.groupingResult.groups.flatMap(g=>g.members).every(s=>s.ho_ten.startsWith('Phiên riêng'))")
                assert first.evaluate("async()=>(await import('/js/state.js')).state.groupingResult.k") == 6
                print("PASS: two contexts isolated; shared-address cache checks each student's gender", flush=True)

                first.locator("#location-filter").select_option("all")
                card = first.locator(".location-card").first
                card.locator(".location-address").fill("Địa chỉ không xác định")
                expect(first.locator("#export-groups")).to_be_disabled()
                assert first.evaluate("async()=>{let s=(await import('/js/state.js')).state;return s.students[0].lat===null&&s.students[0].lon===null&&s.groupingResult===null}")
                card.locator(".location-retry").click()
                expect(first.locator("#geocode-all")).to_be_enabled()
                expect(first.locator("#generate-groups")).to_be_disabled()
                card.locator(".location-address").fill("Cư xá A, " + NLU)
                card.locator(".location-retry").click()
                expect(first.locator("#location-progress")).to_contain_text("30/30")
                assert first.evaluate("async()=>(await import('/js/state.js')).state.students.map(s=>[s.student_id,s.ho_ten,s.gioi_tinh,s.dia_chi_goc])") == before
                first.locator("#generate-groups").click()
                expect(first.locator("#grouping-status")).to_have_attribute("data-state", "success")
                with first.expect_download() as download:
                    first.locator("#export-groups").click()
                download.value.save_as(OUT / "offline-corrected.xlsx")
                print("PASS: address error/edit/retry clears coordinates/groups; identities preserved; export recalculated", flush=True)

                for page in pages:
                    assert page.evaluate("localStorage.length===0&&sessionStorage.length===0")
                    page.reload(wait_until="domcontentloaded")
                    expect(page.locator("#student-rows tr")).to_have_count(0)
                    assert page.evaluate("async()=>{let s=(await import('/js/state.js')).state;return s.students.length===0&&s.groupingResult===null}")
                assert not errors, errors
                (OUT / "offline-summary.json").write_text(json.dumps({"sample_rows": 30, "lookup_calls": len(calls), "initial_lookup_calls": 9, "contexts": 2, "page_errors": errors, "source": "synthetic Geoapify MockTransport"}, indent=2), encoding="utf8")
                browser.close()
                print("PASS: reload clears both sessions; no browser storage or JavaScript exceptions", flush=True)
        finally:
            for process in reversed(processes):
                if process.poll() is None:
                    process.terminate()
                    process.wait(timeout=10)


if __name__ == "__main__":
    main()

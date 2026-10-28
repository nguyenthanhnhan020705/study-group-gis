"""Validate the existing partition and write only in-memory, formula-safe XLSX."""
from io import BytesIO
from math import isclose

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.schemas.export import ExportGroupsRequest
from app.services.metrics_service import aggregate_metrics, pair_metrics


class ExportProblem(ValueError):
    pass


def _check_metrics(actual, expected):
    return actual.pair_count == expected["pair_count"] and all(
        isclose(getattr(actual, field), expected[field], rel_tol=1e-9, abs_tol=1e-6)
        for field in ("mean_distance_m", "max_distance_m")
    )


def validate_export(payload: ExportGroupsRequest) -> None:
    def reject():
        raise ExportProblem("Kết quả chia nhóm không nhất quán. Vui lòng chia nhóm lại trước khi xuất.")

    groups = payload.groups
    n = sum(len(group.members) for group in groups)
    target = payload.target_male + payload.target_female
    if not target or not 0 < n <= 1000 or payload.k != max(1, n // target) or len(groups) != payload.k:
        reject()
    if sorted(group.group_id for group in groups) != list(range(1, payload.k + 1)):
        reject()
    if max(group.size for group in groups) - min(group.size for group in groups) > 1:
        reject()
    ids, metrics = {}, []
    for group in groups:
        male = sum(row.gioi_tinh == "Nam" for row in group.members)
        female = len(group.members) - male
        if (group.size, group.male, group.female, group.male_delta, group.female_delta) != (
            len(group.members), male, female, male - payload.target_male, female - payload.target_female
        ):
            reject()
        for row in group.members:
            if row.student_id in ids:
                reject()
            ids[row.student_id] = group.group_id
            if any(ILLEGAL_CHARACTERS_RE.search(text) for text in (row.ho_ten, row.vi_tri_dai_dien)):
                reject()
        expected = pair_metrics([(row.lat, row.lon) for row in group.members])
        if not _check_metrics(group.metrics, expected):
            reject()
        metrics.append(expected)
    assigned = {}
    for item in payload.assignments:
        if item.student_id in assigned:
            reject()
        assigned[item.student_id] = item.group_id
    if ids != assigned or not _check_metrics(payload.overall_metrics, aggregate_metrics(metrics)):
        reject()


def export_workbook(payload: ExportGroupsRequest) -> bytes:
    validate_export(payload)
    book = Workbook()
    members = book.active
    members.title = "Danh_sach_nhom"
    stats = book.create_sheet("Thong_ke_nhom")
    members.append(["STT", "Nhóm", "Họ và tên", "Giới tính", "Nơi ở để học tập"])
    stats.append(["Nhóm", "Tổng sinh viên", "Số nam", "Số nữ", "Khoảng cách trung bình (m)", "Khoảng cách lớn nhất (m)", "Ghi chú"])
    index = 0
    for group in sorted(payload.groups, key=lambda group: group.group_id):
        for row in sorted(group.members, key=lambda row: row.student_id):
            index += 1
            members.append([index, group.group_id, row.ho_ten, row.gioi_tinh, row.vi_tri_dai_dien])
            for cell in (members.cell(index + 1, 3), members.cell(index + 1, 4), members.cell(index + 1, 5)):
                # Preserve exact text, including a leading '='; an apostrophe changes the value.
                cell.data_type = "s"
        note = "Đạt mục tiêu Nam/Nữ."
        if group.male_delta or group.female_delta:
            note = (f"Mục tiêu {payload.target_male} Nam / {payload.target_female} Nữ; "
                    f"lệch {group.male_delta:+d} Nam / {group.female_delta:+d} Nữ.")
        stats.append([group.group_id, group.size, group.male, group.female,
                      group.metrics.mean_distance_m, group.metrics.max_distance_m, note])
    for sheet, widths in ((members, [8, 10, 32, 12, 90]), (stats, [10, 18, 12, 12, 28, 28, 60])):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17634A")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    for row in stats.iter_rows(min_row=2, min_col=5, max_col=6):
        for cell in row:
            cell.number_format = "0.00"
    output = BytesIO()
    try:
        book.save(output)
        return output.getvalue()
    finally:
        book.close()
        output.close()

"""Generate the public, empty import workbook without adding phantom data rows."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment


def main():
    book = Workbook()
    sheet = book.active
    sheet.title = 'Danh sách sinh viên'
    headers = ['Họ và tên', 'Giới tính', 'Nơi ở để học tập', 'Cư xá/Khu vực']
    sheet.append(headers)
    hints = ['Nhập họ tên đầy đủ.', 'Chọn Nam hoặc Nữ.', 'Nhập địa chỉ nơi ở để học tập, kèm phường/xã và tỉnh/thành phố.', 'Có thể để trống nếu ở nhà riêng. Với KTX ghi rõ cư xá/khu vực.']
    for cell, hint in zip(sheet[1], hints):
        cell.font = Font(bold=True, color='FFFFFF', size=12)
        cell.fill = PatternFill('solid', fgColor='17634A')
        cell.alignment = Alignment(vertical='center')
        cell.comment = Comment(hint, 'WebGIS')
    sheet.row_dimensions[1].height = 30
    for col, width in zip('ABCD', [30, 16, 65, 30]):
        sheet.column_dimensions[col].width = width
    sheet.freeze_panes = 'A2'
    choice = DataValidation(type='list', formula1='"Nam,Nữ"', allow_blank=False)
    choice.errorTitle = 'Giới tính không hợp lệ'
    choice.error = 'Vui lòng chọn Nam hoặc Nữ.'
    choice.showErrorMessage = True
    choice.errorStyle = 'stop'
    sheet.add_data_validation(choice)
    choice.add('B2:B1001')
    guide = book.create_sheet('Hướng dẫn')
    for row in [
        ['HƯỚNG DẪN NHẬP DANH SÁCH'],
        ['1. Điền dữ liệu từ dòng 2 của sheet Danh sách sinh viên; giữ nguyên 4 tiêu đề.'],
        ['2. Ba cột đầu bắt buộc; Cư xá/Khu vực có thể để trống với nhà riêng.'],
        ['3. Giới tính chọn Nam hoặc Nữ; nhập văn bản, không dùng công thức.'],
        ['4. Tối đa 1.000 sinh viên và 5 MiB. Lưu dạng .xlsx rồi tải lên WebGIS.'],
        ['5. Ví dụ bên dưới chỉ để tham khảo, không được hệ thống nhập.'],
        headers,
        ['Nguyễn Văn An', 'Nam', 'KTX Nông Lâm, TP.HCM', 'Cư xá A'],
        ['Trần Thị Bình', 'Nữ', 'Linh Trung, TP.HCM', ''],
    ]:
        guide.append(row)
    guide.column_dimensions['A'].width = 95
    for col in 'BCD':
        guide.column_dimensions[col].width = 30
    path = Path(__file__).resolve().parents[1] / 'frontend/templates/danh-sach-sinh-vien.xlsx'
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


if __name__ == '__main__':
    main()

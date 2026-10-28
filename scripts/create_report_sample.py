"""One reproducible, synthetic presentation workbook with current addresses."""
from pathlib import Path
import sys
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))
from app.services.known_places import NLU_ADDRESS, VNU_B_ADDRESS, CO_MAY_ADDRESS
from app.services.address_labels import STREETS

OUTPUT = ROOT / "data/samples/danh_sach_sinh_vien_bao_cao.xlsx"
VO, KHA = STREETS
CASES = [
    (NLU_ADDRESS, "Cư xá A", ["Nam"] * 3),
    (NLU_ADDRESS, "Cư xá C", ["Nam"] * 3),
    (NLU_ADDRESS, "Cư xá B", ["Nữ"] * 3),
    (NLU_ADDRESS, "Cư xá D", ["Nữ"] * 2),
    (NLU_ADDRESS, "Cư xá E", ["Nam", "Nam", "Nữ", "Nữ"]),
    (CO_MAY_ADDRESS, "KTX Cỏ May", ["Nam", "Nam", "Nữ"]),
    (VNU_B_ADDRESS, "khu B", ["Nam", "Nam", "Nam", "Nữ", "Nữ"]),
    (VO, "Đường Võ Văn Ngân", ["Nam", "Nam", "Nữ"]),
    (KHA, "Đường Kha Vạn Cân", ["Nam", "Nam", "Nam", "Nữ"]),
]
ROWS = [(gender, address, target) for address, target, genders in CASES for gender in genders]


def main():
    book = Workbook()
    sheet = book.active
    sheet.title = "Danh_sach_sinh_vien"
    sheet.append(["Họ và tên", "Giới tính", "Nơi ở để học tập", "Điểm cần xác nhận"])
    info = book.create_sheet("Huong_dan_va_kich_ban")
    info.append(["DỮ LIỆU GIẢ LẬP PHỤC VỤ KIỂM THỬ / BÁO CÁO — không phải sinh viên thật."])
    info.append(["30 sinh viên: 18 Nam, 12 Nữ; 9 điểm đại diện. Không có tọa độ nhập sẵn."])
    info.append(["Upload sheet đầu; ba cột đầu bắt buộc, cột cuối bổ sung tòa/đường để hệ thống đối chiếu tự động."])
    info.append(["Nhấn Đọc danh sách: tự đối chiếu và hiển thị bản đồ; chỉ xử lý các địa chỉ còn mơ hồ hoặc lỗi."])
    info.append(["NLU: A/C dành cho Nam; B/D dành cho Nữ; E/Cỏ May cả hai (quy tắc do chủ dự án cung cấp)."])
    info.append(["Mọi địa chỉ trong mẫu dùng phường và Thành phố Hồ Chí Minh sau sáp nhập."])
    info.append(["Ngoài KTX: điểm đại diện đoạn đường, không thể hiện chính xác một phòng trọ/số nhà."])
    info.append(["Một địa chỉ KTX NLU có nhiều tòa: ghi rõ cư xá trong địa chỉ hoặc cột cuối; không tự đoán từ giới tính."])
    info.append(["Cỏ May là điểm riêng; KTX khu B là điểm đại diện khu, không phải tọa độ từng tòa/phòng."])
    info.append(["Trùng địa chỉ/tọa độ: một marker, popup phải có đủ danh sách sinh viên."])
    info.append(["Kha Vạn Cân có kết quả ở nhiều phường: chỉ chọn kết quả đã đối chiếu phường Thủ Đức."])
    info.append(["Dữ liệu lỗi/thiếu được kiểm thử riêng, không đưa vào danh sách trình bày hợp lệ này."])
    info.append(["Kịch bản / điểm", "Nam", "Nữ", "Tổng trong popup"])
    for address, target, genders in CASES:
        counts = Counter(genders)
        info.append([target, counts["Nam"], counts["Nữ"], len(genders)])
    for i, (gender, address, target) in enumerate(ROWS, 1):
        sheet.append([f"Sinh viên giả lập {i:02d}", gender, address, target])
    sources = book.create_sheet("Nguon_dia_chi")
    sources.append(["Địa điểm", "Nguồn đối chiếu tên địa chỉ", "Lưu ý"])
    from app.services.known_places import resolve_place
    for address in (NLU_ADDRESS, CO_MAY_ADDRESS, VNU_B_ADDRESS):
        sources.append([address, resolve_place(address).address_source_url, "Tọa độ lấy trực tiếp từ Geoapify Places/OSM; kiểm tra rồi xác nhận."])
    for address, (_, _, source) in STREETS.items():
        sources.append([address, source, "Nhãn hiện hành chỉ áp dụng khi khớp đường, phường, vùng và mức tin cậy."])
    for ws in book:
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17634A")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column in ("A", "B", "C", "D"):
            ws.column_dimensions[column].width = 40
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 14
    sheet.column_dimensions["C"].width = 95
    sheet.auto_filter.ref = sheet.dimensions
    for row in range(2, 32):
        sheet.row_dimensions[row].height = 45
    info.column_dimensions["A"].width = 110
    sources.column_dimensions["A"].width = 85
    sources.column_dimensions["B"].width = 95
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    book.save(OUTPUT)
    book.close()
    print(f"Created {OUTPUT.name}: {len(ROWS)} synthetic students / {len(CASES)} points")


if __name__ == "__main__":
    main()
